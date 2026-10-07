#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""PPTX 深挖提取 —— 补 `extract.py` 在课件上的三处硬伤。

`extract.py` 走 python-pptx 的 shape 遍历，遇到课件会漏三类内容：

1. **原生表格抽不到**：课件里的各种表格、数据表常被包在
   `<mc:AlternateContent>` 里，python-pptx 的 `shape.has_table` 会返回 False，
   表格直接消失（一份课件可能有十几页全中）。
   → 本脚本直接解析 `slideN.xml` 的 `<a:tbl>`，绕开 python-pptx。
2. **公式列是空的**：表格里「表示形式」一列、以及整页的推导步骤都写成 OMML
   (`<m:oMath>`)，`<a:t>` 里一个字都没有。
   → 本脚本按页导出 `<m:t>` 的线性文本，能把 `p∧q`、`FOIL_Gain=...` 捞回来。
3. **内嵌图片无页码归属**：`--extract-images` 导出的图没有页号，不知道哪张属于哪页。
   → 本脚本按 `s{页号}_{序号}.png` 命名，配 `_index.txt`，可直接逐个查看。

用法：
  python pptx_deep.py <文件.pptx> <输出目录> [--tag ch2]

产出（都在输出目录下）：
  {tag}_tables.md    所有表格（含 AlternateContent 里的）
  {tag}_formulas.md  所有 OMML 公式，按页分组
  {tag}_shapes.md    形状层文本（递归组合形状）+ 图片清单
  img/{tag}_sNN_M.*  按页编号的内嵌图片 + img/_{tag}_index.txt

注意：表格里公式那几格仍会是空的 —— 那是 OMML，去 {tag}_formulas.md 里按页号对照补。
"""
import argparse
import os
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

try:  # Windows 控制台默认 cp936，中文输出会乱码
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def slide_list(z):
    return sorted(
        (int(re.search(r"slide(\d+)\.xml$", n).group(1)), n)
        for n in z.namelist()
        if re.match(r"ppt/slides/slide\d+\.xml$", n)
    )


def _md_table(rows: list[list[str]]) -> list[str]:
    """把行列表渲染成能直接贴进笔记的 Markdown 表格。

    两处细节和 `extract.py` 的 `_table_to_md` 保持一致：
    补 `| --- |` 分隔行（缺了它 Obsidian 只会显示一堆竖线），
    以及转义单元格里本来就有的 `|`（不转义会被数成多一列，lint 报列数不一致）。
    """
    if not rows:
        return []
    clean = [[c.replace("|", "\\|") for c in r] for r in rows]
    width = max(len(r) for r in clean)
    clean = [r + [""] * (width - len(r)) for r in clean]     # 合并单元格会参差不齐
    out = ["| " + " | ".join(clean[0]) + " |",
           "| " + " | ".join(["---"] * width) + " |"]
    out += ["| " + " | ".join(r) + " |" for r in clean[1:]]
    return out


def dump_tables(z, out):
    lines = []
    for idx, name in slide_list(z):
        root = ET.fromstring(z.read(name))
        for t in root.iter(f"{A}tbl"):
            rows = []
            for tr in t.findall(f"{A}tr"):
                cells = []
                for tc in tr.findall(f"{A}tc"):
                    txt = "".join(x.text or "" for x in tc.iter(f"{A}t"))
                    cells.append(txt.replace("\n", " ").strip())
                rows.append(cells)
            if rows:
                lines.append(f"\n<!-- slide {idx} 表格 -->")
                lines.extend(_md_table(rows))
    _write(out, "\n".join(lines))


def dump_formulas(z, out):
    lines = []
    for idx, name in slide_list(z):
        root = ET.fromstring(z.read(name))
        fs = []
        for om in root.iter(f"{M}oMath"):
            txt = "".join(t.text or "" for t in om.iter(f"{M}t")).strip()
            if txt:
                fs.append(txt)
        if fs:
            lines.append(f"\n<!-- slide {idx} 公式 ({len(fs)}) -->")
            lines.extend("  " + f for f in fs)
    _write(out, "\n".join(lines))


def dump_shapes(path, out):
    from pptx import Presentation
    prs = Presentation(path)
    lines = []

    def walk(shapes, res):
        for sh in shapes:
            if sh.__class__.__name__ == "GroupShape":
                walk(sh.shapes, res)
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                res.append(("TABLE", sh.table))
                continue
            if sh.__class__.__name__ == "Picture":
                res.append(("PIC", sh.name))
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    res.append(("TXT", t))
                continue
            res.append(("SHAPE", str(sh.shape_type)))

    for i, slide in enumerate(prs.slides, 1):
        lines.append(f"\n<!-- ===== slide {i} ===== -->")
        items = []
        try:
            walk(slide.shapes, items)
        except Exception as e:                                    # noqa: BLE001
            lines.append(f"[ERR] {e}")
            continue
        for kind, val in items:
            if kind == "TABLE":
                lines.append("**TABLE**")
                lines.extend(_md_table([
                    [c.text.replace("\n", " ").strip() for c in row.cells]
                    for row in val.rows]))
                lines.append("")
            elif kind == "TXT":
                lines.append("TXT: " + val.replace("\n", "\n     "))
            elif kind == "PIC":
                lines.append("PIC: " + val)
            else:
                lines.append(f"{kind}: {val}")
    _write(out, "\n".join(lines))


def dump_pics(path, imgdir, tag):
    from pptx import Presentation
    os.makedirs(imgdir, exist_ok=True)
    prs = Presentation(path)
    rows = []
    for i, slide in enumerate(prs.slides, 1):
        n = 0
        stack = [slide.shapes]
        # 用显式栈代替递归，避免闭包里的 nonlocal 计数踩坑
        while stack:
            shapes = stack.pop()
            for sh in shapes:
                if sh.__class__.__name__ == "GroupShape":
                    stack.append(sh.shapes)
                    continue
                if sh.__class__.__name__ == "Picture":
                    n += 1
                    img = sh.image
                    fn = f"{tag}_s{i:02d}_{n}.{img.ext}"
                    with open(os.path.join(imgdir, fn), "wb") as f:
                        f.write(img.blob)
                    rows.append((i, fn, len(img.blob)))
    _write(os.path.join(imgdir, f"_{tag}_index.txt"),
           "\n".join(f"slide {i:02d}  {fn}  {sz/1024:.0f} KB" for i, fn, sz in rows))
    print(f"[图片] {len(rows)} 张 → {imgdir}")


def _write(path, text):
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text + "\n")
    print(f"[完成] {path}  ({len(text)} 字符)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("outdir")
    ap.add_argument("--tag", default="ppt")
    a = ap.parse_args()

    os.makedirs(a.outdir, exist_ok=True)
    z = zipfile.ZipFile(a.pptx)
    dump_tables(z, os.path.join(a.outdir, f"{a.tag}_tables.md"))
    dump_formulas(z, os.path.join(a.outdir, f"{a.tag}_formulas.md"))
    dump_shapes(a.pptx, os.path.join(a.outdir, f"{a.tag}_shapes.md"))
    dump_pics(a.pptx, os.path.join(a.outdir, "img"), a.tag)


if __name__ == "__main__":
    sys.exit(main())
